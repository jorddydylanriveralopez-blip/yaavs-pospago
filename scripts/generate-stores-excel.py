#!/usr/bin/env python3
"""Generate the polished Excel directory of AT&T store landings."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
STORES_JS = ROOT / "js" / "tiendas-att-stores.js"
BASE = "https://yaavs-pospago.hostingersite.com"
CENTRAL_WA = "525522331210"

INK, CYAN, BLUE, WHITE = "0A2540", "00C1D4", "009FDB", "FFFFFF"
SOFT, GREY, LINE, GREEN, FB = "E8F6FB", "5A7388", "C5D8E4", "25D366", "1877F2"


def load_stores() -> list[dict]:
    text = STORES_JS.read_text(encoding="utf-8")
    m = re.search(r"window\.YAAVS_ATT_STORES\s*=\s*(\[.*\])\s*;", text, re.S)
    if not m:
        raise SystemExit("Could not parse stores")
    return json.loads(m.group(1))


def title_case(s: str) -> str:
    return re.sub(r"(^|[\s([])\S", lambda m: m.group(0).upper(), (s or "").lower())


def fmt_phone(raw: str) -> tuple[str, str]:
    d = re.sub(r"\D", "", str(raw or ""))
    if d.startswith("52") and len(d) == 12:
        local = d[2:]
        return f"+52 {local[0:2]} {local[2:6]} {local[6:]}", f"https://wa.me/{d}"
    if len(d) == 10:
        return f"+52 {d[0:2]} {d[2:6]} {d[6:]}", f"https://wa.me/52{d}"
    if d:
        return f"+{d}", f"https://wa.me/{d}"
    return "—", ""


def build(out: Path) -> None:
    stores_raw = load_stores()
    stores = sorted(stores_raw, key=lambda s: (s.get("state") or "", s.get("name") or ""))
    thin = Border(
        left=Side(style="thin", color=LINE),
        right=Side(style="thin", color=LINE),
        top=Side(style="thin", color=LINE),
        bottom=Side(style="thin", color=LINE),
    )
    wb = Workbook()

    cover = wb.active
    cover.title = "Portada"
    cover.sheet_view.showGridLines = False
    for col, w in enumerate([4, 28, 22, 22, 22, 18, 4], 1):
        cover.column_dimensions[get_column_letter(col)].width = w

    cover.merge_cells("B2:F2")
    cover["B2"] = "YAAVS POSPAGO  ·  AT&T"
    cover["B2"].font = Font(name="Calibri", size=14, bold=True, color=CYAN)

    cover.merge_cells("B3:F3")
    cover["B3"] = "Directorio de sucursales"
    cover["B3"].font = Font(name="Calibri", size=28, bold=True, color=INK)

    cover.merge_cells("B4:F4")
    cover["B4"] = "33 landings · redes · ubicación · contacto"
    cover["B4"].font = Font(name="Calibri", size=13, color=GREY)

    cover.merge_cells("B6:F6")
    cover["B6"] = f"Actualizado: {date.today().strftime('%d/%m/%Y')}   ·   Sitio: {BASE}"
    cover["B6"].font = Font(name="Calibri", size=11, color=GREY)

    cover.merge_cells("B8:C9")
    cover["B8"] = "33"
    cover["B8"].font = Font(name="Calibri", size=36, bold=True, color=WHITE)
    cover["B8"].fill = PatternFill("solid", fgColor=INK)
    cover["B8"].alignment = Alignment(horizontal="center", vertical="center")
    cover.merge_cells("B10:C10")
    cover["B10"] = "SUCURSALES"
    cover["B10"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    cover["B10"].fill = PatternFill("solid", fgColor=SOFT)
    cover["B10"].alignment = Alignment(horizontal="center")

    cover.merge_cells("D8:E9")
    cover["D8"] = str(len({s["state"] for s in stores}))
    cover["D8"].font = Font(name="Calibri", size=36, bold=True, color=WHITE)
    cover["D8"].fill = PatternFill("solid", fgColor=BLUE)
    cover["D8"].alignment = Alignment(horizontal="center", vertical="center")
    cover.merge_cells("D10:E10")
    cover["D10"] = "ESTADOS"
    cover["D10"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    cover["D10"].fill = PatternFill("solid", fgColor=SOFT)
    cover["D10"].alignment = Alignment(horizontal="center")

    cover.merge_cells("F8:F9")
    cover["F8"] = "WA"
    cover["F8"].font = Font(name="Calibri", size=28, bold=True, color=WHITE)
    cover["F8"].fill = PatternFill("solid", fgColor=GREEN)
    cover["F8"].alignment = Alignment(horizontal="center", vertical="center")
    cover["F10"] = "CENTRAL"
    cover["F10"].font = Font(name="Calibri", size=11, bold=True, color=INK)
    cover["F10"].fill = PatternFill("solid", fgColor=SOFT)
    cover["F10"].alignment = Alignment(horizontal="center")

    cover.merge_cells("B12:F12")
    cover["B12"] = "WhatsApp central de cotización"
    cover["B12"].font = Font(name="Calibri", size=12, bold=True, color=INK)
    cover.merge_cells("B13:F13")
    cover["B13"] = "+52 55 2233 1210"
    cover["B13"].font = Font(name="Calibri", size=18, bold=True, color=GREEN)
    cover["B13"].hyperlink = f"https://wa.me/{CENTRAL_WA}"

    cover.merge_cells("B15:F15")
    cover["B15"] = "Cómo usar este archivo"
    cover["B15"].font = Font(name="Calibri", size=13, bold=True, color=INK)
    for i, tip in enumerate(
        [
            "1. Ve a la hoja «Directorio» — ahí están las 33 sucursales con filtros.",
            "2. Haz clic en Landing / Facebook / Maps / WhatsApp para abrir el enlace.",
            "3. La hoja «Por estado» agrupa las tiendas para compartir por zona.",
            "4. «Leyenda» explica cada columna y el formato de los links.",
        ]
    ):
        r = 16 + i
        cover.merge_cells(f"B{r}:F{r}")
        cover[f"B{r}"] = tip
        cover[f"B{r}"].font = Font(name="Calibri", size=11, color=GREY)

    cover.row_dimensions[3].height = 40
    cover.row_dimensions[8].height = 28
    cover.row_dimensions[9].height = 28

    ws = wb.create_sheet("Directorio", 1)
    headers = [
        "#",
        "Sucursal",
        "Ciudad",
        "Estado",
        "Landing (página web)",
        "Dirección",
        "Horario",
        "Google Maps",
        "Facebook",
        "WhatsApp sucursal",
        "Tel. gerente",
        "Gerente",
        "Lat",
        "Lng",
        "ID",
    ]
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:O{len(stores)+1}"
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[1].height = 36

    for i, s in enumerate(stores, 1):
        r = i + 1
        phone_disp, wa_url = fmt_phone(s.get("managerPhone"))
        page_url = f"{BASE}/{s['page']}"
        fb = (s.get("facebook") or "").strip()
        maps = (s.get("mapsLink") or "").strip()
        if not maps and s.get("lat") is not None:
            maps = f"https://www.google.com/maps?q={s['lat']},{s['lng']}"
        row = [
            i,
            s["name"],
            title_case(s.get("city") or ""),
            title_case(s.get("state") or ""),
            page_url,
            s.get("address") or "",
            s.get("hours") or "",
            maps or "—",
            fb or "Sin Facebook (reubicación)",
            wa_url or "—",
            phone_disp,
            s.get("manager") or "",
            s.get("lat"),
            s.get("lng"),
            s.get("id") or "",
        ]
        fill = PatternFill("solid", fgColor=SOFT if i % 2 else WHITE)
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c, val)
            cell.font = Font(name="Calibri", size=10, color=INK)
            cell.fill = fill
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=(c in (5, 6, 7, 9)))
        ws.cell(r, 5).hyperlink = page_url
        ws.cell(r, 5).font = Font(name="Calibri", size=10, color=BLUE, underline="single")
        if maps and maps != "—":
            ws.cell(r, 8).hyperlink = maps
            ws.cell(r, 8).font = Font(name="Calibri", size=10, color=CYAN, underline="single")
            ws.cell(r, 8).value = "Abrir Maps"
        if fb:
            ws.cell(r, 9).hyperlink = fb
            ws.cell(r, 9).font = Font(name="Calibri", size=10, color=FB, underline="single")
            ws.cell(r, 9).value = "Facebook"
        else:
            ws.cell(r, 9).font = Font(name="Calibri", size=10, italic=True, color=GREY)
        if wa_url:
            ws.cell(r, 10).hyperlink = wa_url
            ws.cell(r, 10).font = Font(name="Calibri", size=10, color="1EBE57", underline="single")
            ws.cell(r, 10).value = "WhatsApp"
        ws.row_dimensions[r].height = 42

    for col, w in {
        "A": 5, "B": 28, "C": 22, "D": 18, "E": 42, "F": 55, "G": 38,
        "H": 14, "I": 14, "J": 16, "K": 18, "L": 32, "M": 12, "N": 12, "O": 22,
    }.items():
        ws.column_dimensions[col].width = w

    by = wb.create_sheet("Por estado", 2)
    by.sheet_view.showGridLines = False
    by["A1"] = "Sucursales por estado"
    by["A1"].font = Font(name="Calibri", size=18, bold=True, color=INK)
    by.merge_cells("A1:F1")
    by["A2"] = "Agrupadas para compartir por zona · links listos para copiar"
    by["A2"].font = Font(name="Calibri", size=11, color=GREY)
    by.merge_cells("A2:F2")

    row = 4
    groups: dict[str, list] = defaultdict(list)
    for s in stores:
        groups[s.get("state") or "SIN ESTADO"].append(s)
    for state in sorted(groups.keys()):
        items = groups[state]
        by.merge_cells(f"A{row}:F{row}")
        by[f"A{row}"] = f"{title_case(state)}  ·  {len(items)} sucursal{'es' if len(items) != 1 else ''}"
        by[f"A{row}"].font = Font(name="Calibri", size=12, bold=True, color=WHITE)
        by[f"A{row}"].fill = PatternFill("solid", fgColor=BLUE)
        by.row_dimensions[row].height = 26
        row += 1
        for c, h in enumerate(["Sucursal", "Ciudad", "Landing", "Maps", "Facebook", "WhatsApp"], 1):
            cell = by.cell(row, c, h)
            cell.font = Font(name="Calibri", size=10, bold=True, color=INK)
            cell.fill = PatternFill("solid", fgColor=SOFT)
            cell.border = thin
        row += 1
        for s in sorted(items, key=lambda x: x["name"]):
            page_url = f"{BASE}/{s['page']}"
            _, wa_url = fmt_phone(s.get("managerPhone"))
            fb = (s.get("facebook") or "").strip()
            maps = (s.get("mapsLink") or "").strip() or "—"
            vals = [s["name"], title_case(s.get("city") or ""), page_url, maps, fb or "—", wa_url or "—"]
            for c, v in enumerate(vals, 1):
                cell = by.cell(row, c, v)
                cell.font = Font(name="Calibri", size=10, color=INK)
                cell.border = thin
            by.cell(row, 3).hyperlink = page_url
            by.cell(row, 3).font = Font(name="Calibri", size=10, color=BLUE, underline="single")
            by.cell(row, 3).value = "Abrir landing"
            if maps != "—":
                by.cell(row, 4).hyperlink = maps
                by.cell(row, 4).font = Font(name="Calibri", size=10, color=CYAN, underline="single")
                by.cell(row, 4).value = "Maps"
            if fb:
                by.cell(row, 5).hyperlink = fb
                by.cell(row, 5).font = Font(name="Calibri", size=10, color=FB, underline="single")
                by.cell(row, 5).value = "Facebook"
            if wa_url:
                by.cell(row, 6).hyperlink = wa_url
                by.cell(row, 6).font = Font(name="Calibri", size=10, color="1EBE57", underline="single")
                by.cell(row, 6).value = "WhatsApp"
            by.row_dimensions[row].height = 22
            row += 1
        row += 1
    for col, w in zip("ABCDEF", [28, 22, 16, 12, 12, 12]):
        by.column_dimensions[col].width = w

    leg = wb.create_sheet("Leyenda", 3)
    leg.sheet_view.showGridLines = False
    leg["A1"] = "Leyenda y notas"
    leg["A1"].font = Font(name="Calibri", size=18, bold=True, color=INK)
    leg.merge_cells("A1:B1")
    notes = [
        ("Landing", "Página web propia de la sucursal (tienda/slug.html). Abre con un clic."),
        ("Google Maps", "Cómo llegar a la fachada / ubicación de la tienda."),
        ("Facebook", "Página o perfil de la sucursal. Saltillo 400 no tiene FB (reubicación)."),
        ("WhatsApp sucursal", "Contacto del gerente / tienda (número local)."),
        ("Tel. gerente", "Mismo número, formato legible (+52 …)."),
        ("WhatsApp central", "55 2233 1210 — cotizaciones generales YAAVS Pospago."),
        ("Filtros", "En «Directorio» usa la fila de filtros para buscar por estado o ciudad."),
    ]
    leg["A3"] = "Campo"
    leg["B3"] = "Qué significa"
    for c in (1, 2):
        leg.cell(3, c).font = Font(name="Calibri", bold=True, color=WHITE)
        leg.cell(3, c).fill = PatternFill("solid", fgColor=INK)
        leg.cell(3, c).border = thin
    for i, (a, b) in enumerate(notes, 4):
        leg.cell(i, 1, a).font = Font(name="Calibri", bold=True, color=BLUE)
        leg.cell(i, 2, b).font = Font(name="Calibri", color=INK)
        for c in (1, 2):
            leg.cell(i, c).border = thin
            leg.cell(i, c).alignment = Alignment(wrap_text=True, vertical="center")
        leg.row_dimensions[i].height = 32
    leg.column_dimensions["A"].width = 22
    leg.column_dimensions["B"].width = 78

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out} ({out.stat().st_size} bytes)")


if __name__ == "__main__":
    build(ROOT / "YAAVS-Directorio-Sucursales-ATT.xlsx")
    art = Path("/opt/cursor/artifacts/YAAVS-Directorio-Sucursales-AT&T.xlsx")
    build(art)
